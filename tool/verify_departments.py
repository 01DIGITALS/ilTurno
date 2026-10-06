"""Inspect the Android-generated department PDF/XLSX fixtures without changing them."""
from pathlib import Path
import os
import subprocess
import json
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
output = root / "build/qa/departments-exports"
output.mkdir(parents=True, exist_ok=True)
adb = Path(os.environ["LOCALAPPDATA"]) / "Android/Sdk/platform-tools/adb.exe"
for name in ["departments.pdf", "departments.xlsx"]:
    with (output / name).open("wb") as stream:
        subprocess.run([str(adb), "-s", "emulator-5554", "exec-out", "run-as", "it.sanges.ilturno", "cat", f"files/qa/{name}"], stdout=stream, check=True)
pdf = PdfReader(output / "departments.pdf")
text = " ".join(page.extract_text() for page in pdf.pages)
assert len(pdf.pages) == 1
assert text.count("Cucina") == 1
assert text.count("Sala") == 1
assert text.count("Anna") == 2
sheet = load_workbook(output / "departments.xlsx")["Turni"]
assert sheet.max_row == 16 and sheet.max_column == 5
assert sheet["E2"].value == "Reparto"
assert all(sheet.cell(row, 5).value == "Cucina" for row in range(3, 10))
assert all(sheet.cell(row, 5).value == "Sala" for row in range(10, 17))
assert sheet["D4"].value == "Anna" and sheet["C11"].value == "Anna"
assert sheet["B3"].is_date and sheet["B3"].value.strftime("%Y-%m-%d") == "2026-10-05"
assert sheet["B16"].value.strftime("%Y-%m-%d") == "2026-10-11"
report = {"pdf_pages": len(pdf.pages), "department_labels": True, "xlsx_rows": sheet.max_row, "xlsx_typed_dates": True}
(output / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
subprocess.run(["pdftoppm", "-scale-to", "1200", "-png", str(output / "departments.pdf"), str(output / "departments")], check=True)
print(json.dumps(report))
