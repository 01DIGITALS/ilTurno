"""Read and render the Android-generated QA files; never rewrites an export."""
import os
from pathlib import Path
import subprocess
import json
import unicodedata
import re
import argparse
from pypdf import PdfReader
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--serial', default='emulator-5554')
args = parser.parse_args()
out = root / 'build' / 'qa' / ('exports' if args.serial == 'emulator-5554' else 'exports-android17')
out.mkdir(parents=True, exist_ok=True)
adb = Path(os.environ['LOCALAPPDATA']) / 'Android/Sdk/platform-tools/adb.exe'
files = ['Turni_2026-10-05_2026-10-11.pdf', 'Turni_2026-10-05_2026-10-11.xlsx', 'long-names.pdf']
for name in files:
    with (out / name).open('wb') as stream:
        subprocess.run([str(adb), '-s', args.serial, 'exec-out', 'run-as', 'it.sanges.ilturno', 'cat', f'files/qa/{name}'], stdout=stream, check=True)
pdf = PdfReader(out / files[0])
text = ' '.join(pdf.pages[0].extract_text().split())
for expected in ['TURNI', '5 – 11 ottobre 2026', 'Pranzo', 'Cena', 'lun 5', 'mar 6', 'dom 11', 'Marco', 'Anna']:
    assert expected in text, (expected, text)
assert len(pdf.pages) == 1
assert tuple(map(float, (pdf.pages[0].mediabox.width, pdf.pages[0].mediabox.height))) == (842.0, 595.0)
xlsx = load_workbook(out / files[1])
sheet = xlsx['Turni']
assert sheet['A1'].value == 'TURNI · 5 – 11 ottobre 2026'
assert [sheet.cell(2, column).value for column in range(1, 5)] == ['Giorno', 'Data', 'Pranzo', 'Cena']
assert sheet.max_row == 9
assert sheet['B3'].value.strftime('%Y-%m-%d') == '2026-10-05'
assert sheet['B9'].value.strftime('%Y-%m-%d') == '2026-10-11'
assert sheet['C4'].value == 'Anna, Marco'
assert sheet['D4'].value == 'Marco'
assert sheet['B4'].is_date and sheet['B4'].number_format == 'dd/mm/yyyy'
assert sheet.freeze_panes == 'A3'
long_pdf = PdfReader(out / files[2])
long_text = ' '.join(unicodedata.normalize('NFKC', ' '.join(page.extract_text() for page in long_pdf.pages)).split())
# Repeated page headers/footers sit between the two halves of a wrapped name.
long_text = long_text.replace('TURNI 5 – 11 ottobre 2026 Giorno Pranzo Cena', ' ')
long_text = re.sub(r'ilTurno · Pagina \d+|(?:lun|mar|mer|gio|ven|sab|dom) \d+|Nessun turno', ' ', long_text)
long_text = ' '.join(long_text.split())
counts = {str(i): long_text.count(f'Persona {i} con un nome lungo per la verifica della stampa') for i in range(1, 81)}
assert set(counts.values()) == {7}, counts
report = {'pdf_pages': len(pdf.pages), 'pdf_A4': True, 'pdf_period_correct': True, 'xlsx_typed_dates': True,
          'xlsx_rows': sheet.max_row, 'long_pdf_pages': len(long_pdf.pages), 'long_name_occurrences': counts}
(out / 'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
subprocess.run(['pdftoppm', '-scale-to', '1500', '-png', str(out / files[0]), str(out / 'week')], check=True)
subprocess.run(['pdftoppm', '-scale-to', '900', '-png', str(out / files[2]), str(out / 'long')], check=True)
print(json.dumps({key: value for key, value in report.items() if key != 'long_name_occurrences'}, indent=2))
print('Long-name counts:', sorted(set(counts.values())))
